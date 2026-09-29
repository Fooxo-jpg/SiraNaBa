from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "deliverables"
ASSET_DIR = ROOT / "tools" / "manual_assets"
OUT_DIR.mkdir(exist_ok=True)
ASSET_DIR.mkdir(exist_ok=True)
OUT = OUT_DIR / "SiraNaBa_User_Manual_Small_Team.docx"

GREEN = "2F6B4F"
PALE_GREEN = "EAF4EE"
SAND = "F7F4EC"
INK = "1F2933"
MID = "52606D"
LINE = "D9D9D9"
WHITE = "FFFFFF"


def image_font(size, bold=False):
    names = ["aptos-bold.ttf", "calibrib.ttf", "arialbd.ttf"] if bold else ["aptos.ttf", "calibri.ttf", "arial.ttf"]
    for name in names:
        path = Path("C:/Windows/Fonts") / name
        if path.exists():
            return ImageFont.truetype(path, size)
    return ImageFont.load_default()


def rounded(draw, xy, radius, fill, outline=LINE, width=2):
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)


def create_portal_wireframe(path):
    im = Image.new("RGB", (1600, 670), "#F7F4EC")
    d = ImageDraw.Draw(im)
    small = image_font(22)
    label = image_font(23, True)
    title = image_font(34, True)
    d.text((55, 35), "SiraNaBa at a glance", font=title, fill="#1F2933")
    cards = [
        (55, 105, 500, 595, "Finance Management", "PHP 8,450.00", ["Rent balance", "Utility statements", "Payment history"], "Pay Now"),
        (575, 105, 1020, 595, "AI Ticket Triage", "HIGH PRIORITY", ["Request received", "Urgency checked", "Staff assignment"], "View Ticket"),
        (1095, 105, 1540, 595, "Notifications", "3 UNREAD", ["Payment reminder", "Ticket assigned", "Facility notice"], "Open Alerts"),
    ]
    for x1, y1, x2, y2, heading, metric, rows, button in cards:
        rounded(d, (x1, y1, x2, y2), 24, "white")
        d.text((x1 + 30, y1 + 28), heading, font=label, fill="#1F2933")
        d.text((x1 + 30, y1 + 85), metric, font=title, fill="#2F6B4F")
        y = y1 + 165
        for row in rows:
            d.ellipse((x1 + 32, y + 8, x1 + 46, y + 22), fill="#2F6B4F")
            d.text((x1 + 62, y), row, font=small, fill="#52606D")
            y += 65
        rounded(d, (x1 + 30, y2 - 90, x2 - 30, y2 - 30), 13, "#2F6B4F", "#2F6B4F")
        tw = d.textbbox((0, 0), button, font=label)[2]
        d.text(((x1 + x2 - tw) / 2, y2 - 76), button, font=label, fill="white")
    im.save(path)


def set_font(run, size=9.6, bold=False, color=INK, name="Aptos"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    shade = props.find(qn("w:shd"))
    if shade is None:
        shade = OxmlElement("w:shd")
        props.append(shade)
    shade.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=85, start=105, bottom=85, end=105):
    props = cell._tc.get_or_add_tcPr()
    margins = props.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        props.append(margins)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LINE, size=6):
    props = table._tbl.tblPr
    borders = props.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        props.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:color"), color)


def repeat_header(row):
    props = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    props.append(marker)


def text(doc, value, size=9.6, bold=False, color=INK, before=0, after=4, line=1.06):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    set_font(p.add_run(value), size, bold, color)
    return p


def heading(doc, value, level=1, before=5, after=3):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_with_next = True
    set_font(p.add_run(value), 13.2 if level == 1 else 10.6, True, "000000")
    return p


def step(doc, number, title, body):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2.4)
    p.paragraph_format.line_spacing = 1.02
    set_font(p.add_run(f"{number}  {title}  "), 9.5, True, GREEN)
    set_font(p.add_run(body), 9.5)


def bullet(doc, lead, body):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.15)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.02
    set_font(p.add_run(f"{lead}: "), 9.3, True, GREEN)
    set_font(p.add_run(body), 9.3)


def data_table(doc, headers, rows, widths, font_size=8.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    set_table_borders(table)
    for idx, width in enumerate(widths):
        table.columns[idx].width = Inches(width)
    for idx, value in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, GREEN)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(value), 8.7, True, WHITE)
    repeat_header(table.rows[0])
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        if ridx % 2:
            for cell in cells:
                set_cell_shading(cell, "F4F8F5")
        for cell, value in zip(cells, row):
            set_cell_margins(cell, 68, 95, 68, 95)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            set_font(p.add_run(value), font_size)
    return table


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(paragraph.add_run("SiraNaBa User Manual   |   "), 8, False, MID)
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def build():
    visual = ASSET_DIR / "siranaba_three_features_wireframe.png"
    create_portal_wireframe(visual)

    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.48)
    sec.bottom_margin = Inches(0.45)
    sec.left_margin = Inches(0.6)
    sec.right_margin = Inches(0.6)
    sec.footer_distance = Inches(0.23)
    add_page_number(sec.footer.paragraphs[0])

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.6)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(3)
    for name in ("Title", "Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.name = "Aptos Display" if name == "Title" else "Aptos"
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)

    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_after = Pt(1)
    set_font(title.add_run("SiraNaBa User Manual"), 23, True, "000000", "Aptos Display")
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(5)
    set_font(subtitle.add_run("End user guide for a small team of 5 to 20 users"), 10.7, True, GREEN)
    text(doc, "SiraNaBa is a web-based facility portal for tenants and property staff. It keeps bills and payments, maintenance requests, AI-assisted ticket priority, and service notifications in one place. This guide explains the everyday tasks most users need.", after=5)

    heading(doc, "Getting Started", 1, 3, 2)
    step(doc, "1", "Open the portal", "Use the web address supplied by your administrator in Chrome, Edge, Safari, or another current browser. No installation is required.")
    step(doc, "2", "Sign in", "Enter your registered email address and password, then select Sign In. Administrators open the Admin Sign In page. Your role decides which screens you can see.")
    step(doc, "3", "Use the menu", "Tenants use Dashboard, Maintenance, Billing, and Notifications. Property staff also use the administration area to manage tenants, payments, and ticket dispatch.")
    step(doc, "4", "Finish safely", "Select Sign Out when using a shared device. Ask your administrator for a password reset if you cannot sign in.")

    heading(doc, "Three Main Features", 1, 5, 2)
    bullet(doc, "Finance Management", "Review rent and utility balances, open itemised utility statements, record or simulate a payment, and keep the receipt reference. The project payment screen clearly states when no real funds are transferred.")
    bullet(doc, "Maintenance Tickets and AI Triage", "Submit the problem, location, severity information, and optional photos. SiraNaBa records the ticket and checks its urgency so staff can prioritise and assign the right person. Staff remain responsible for the final decision.")
    bullet(doc, "Notifications", "See payment reminders, ticket progress, scheduled work, and facility notices. Opening an alert takes you to the related item and marks the alert as read.")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(1)
    p.add_run().add_picture(str(visual), width=Inches(7.05))
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(0)
    set_font(caption.add_run("Wireframe showing the three main areas in SiraNaBa"), 8, False, MID)

    doc.add_page_break()
    heading(doc, "What Happens When You Act", 1, 0, 2)
    text(doc, "An event is something you do or something that changes. The action is SiraNaBa's response.", 9.1, after=3)
    event_rows = [
        ("You select Pay Now and confirm", "The system validates the amount, updates the balance, adds the transaction to Payment History, and shows a receipt or failure message."),
        ("You submit a maintenance request", "The request receives a ticket number. AI triage checks the details and suggests a priority for staff review."),
        ("A staff member assigns or updates a ticket", "The ticket timeline changes and a Maintenance notification appears for the tenant."),
        ("A bill, due date, or payment changes", "A Payments notification shows the new information and links to Billing."),
        ("You select a notification", "The related bill, ticket, or notice opens and the notification is marked as read."),
    ]
    data_table(doc, ("Event", "Action"), event_rows, (2.5, 4.65), 8.15)

    heading(doc, "Quick How To", 1, 6, 2)
    step(doc, "Finance", "Open Billing", "Check Total Outstanding and the due dates. Select Pay Now, choose Rent, Utilities, or Both, enter the amount, confirm, and save the reference code shown on the receipt.")
    step(doc, "Ticket", "Open Maintenance", "Select Submit Maintenance Request. Complete Details, Location, Severity, and Review; add a clear photo when useful; then submit and track the status under Maintenance.")
    step(doc, "Alerts", "Open Notifications", "Use All Alerts or a category filter, select an alert to open its related item, and use Mark all as read only after the team has handled the items.")

    heading(doc, "Troubleshooting and Frequently Asked Questions", 1, 5, 2)
    faq_rows = [
        ("Why can I not sign in?", "Check your email and password. Make sure you are using the tenant or admin sign-in page that matches your role. If it still fails, ask the administrator to confirm your account."),
        ("Why did my event not trigger?", "Refresh the page and wait briefly. Confirm the payment or ticket was successfully submitted. Check Notifications and its filters. If the item is still missing, report the time and ticket or receipt reference to the administrator."),
        ("Why is ticket priority still loading?", "The triage check may still be running or may not have enough information. Refresh Maintenance. Staff can review and update the priority if needed."),
        ("Why did a payment fail?", "Read the message shown, confirm the amount does not exceed the outstanding balance, and try once more. Do not repeatedly submit while a payment is marked pending."),
        ("Why can I not see a bill or ticket?", "Clear search and filters, confirm you are signed into the correct account, and select Retry if the page could not load. Contact the administrator if the item belongs to another account."),
    ]
    data_table(doc, ("Question", "What to do"), faq_rows, (2.35, 4.8), 7.75)

    heading(doc, "Small Team Working Rules", 1, 5, 2)
    text(doc, "For a 5 to 20 person team, assign one finance owner to check payments, one maintenance coordinator to review AI priority and dispatch work, and one backup administrator. Use the ticket number or payment reference when discussing an issue, and mark notifications read only after someone has acted.", 8.8, after=0, line=1.02)

    core = doc.core_properties
    core.title = "SiraNaBa User Manual"
    core.subject = "End-user documentation for a small team"
    core.author = "SiraNaBa Project"
    core.keywords = "SiraNaBa, user manual, finance, ticket triage, notifications"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
