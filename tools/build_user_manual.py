from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "deliverables"
ASSET_DIR = ROOT / "tools" / "manual_assets"
OUT_DIR.mkdir(exist_ok=True)
ASSET_DIR.mkdir(exist_ok=True)
OUT = OUT_DIR / "SiraNaBa_User_Manual_Small_Team.docx"

GREEN = "2F6B4F"
GREEN_LIGHT = "EAF4EE"
SAND = "F7F4EC"
INK = "1F2933"
MID = "52606D"
LINE = "D9D9D9"
WHITE = "FFFFFF"


def font(size, bold=False):
    candidates = [
        r"C:\Windows\Fonts\aptos.ttf",
        r"C:\Windows\Fonts\calibri.ttf",
        r"C:\Windows\Fonts\arial.ttf",
    ]
    if bold:
        candidates = [
            r"C:\Windows\Fonts\aptos-bold.ttf",
            r"C:\Windows\Fonts\calibrib.ttf",
            r"C:\Windows\Fonts\arialbd.ttf",
        ] + candidates
    for path in candidates:
        if Path(path).exists():
            return ImageFont.truetype(path, size)
    return ImageFont.load_default()


def rounded(draw, xy, radius, fill, outline=None, width=1):
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)


def toggle(draw, x, y, on=True):
    color = "#2F6B4F" if on else "#CBD2D9"
    rounded(draw, (x, y, x + 64, y + 32), 16, color)
    cx = x + 48 if on else x + 16
    draw.ellipse((cx - 12, y + 4, cx + 12, y + 28), fill="white")


def create_settings_wireframe(path):
    im = Image.new("RGB", (1400, 660), "#F7F4EC")
    d = ImageDraw.Draw(im)
    f12, f16, f20, f28 = font(21), font(25), font(30, True), font(38, True)
    rounded(d, (55, 45, 1345, 615), 28, "white", "#D9D9D9", 3)
    d.text((105, 85), "Alert Settings", font=f28, fill="#1F2933")
    d.text((105, 145), "DELIVERY CHANNELS", font=f12, fill="#52606D")
    rows = [
        ("Email", "A copy of every alert in your inbox", True),
        ("SMS", "Text messages for time-sensitive alerts", False),
        ("Push Notifications", "Alerts on this device in real time", True),
    ]
    y = 188
    for label, desc, on in rows:
        d.text((115, y), label, font=f20, fill="#1F2933")
        d.text((115, y + 38), desc, font=f16, fill="#52606D")
        toggle(d, 1195, y + 12, on)
        y += 90
    d.line((105, 465, 1295, 465), fill="#E4E7EB", width=2)
    d.text((105, 492), "ALERT CATEGORIES", font=f12, fill="#52606D")
    x = 115
    for label in ["Payments", "Maintenance", "Community"]:
        rounded(d, (x, 535, x + 300, 585), 18, "#EAF4EE", "#B7D8C4", 2)
        d.text((x + 22, 545), label, font=f16, fill="#2F6B4F")
        toggle(d, x + 220, 544, True)
        x += 370
    im.save(path)


def create_bug_wireframe(path):
    im = Image.new("RGB", (1400, 510), "#F7F4EC")
    d = ImageDraw.Draw(im)
    f14, f18, f22, f30 = font(22), font(27), font(31, True), font(40, True)
    rounded(d, (55, 35, 1345, 475), 28, "white", "#D9D9D9", 3)
    d.text((105, 70), "Submit Maintenance Request", font=f30, fill="#1F2933")
    steps = [("1", "Details", True), ("2", "Location", False), ("3", "Review", False)]
    x = 130
    for n, label, active in steps:
        fill = "#2F6B4F" if active else "white"
        outline = "#2F6B4F" if active else "#CBD2D9"
        text_fill = "white" if active else "#52606D"
        d.ellipse((x, 138, x + 48, 186), fill=fill, outline=outline, width=3)
        d.text((x + 16, 147), n, font=f14, fill=text_fill)
        d.text((x + 62, 145), label, font=f18, fill="#1F2933")
        if x < 800:
            d.line((x + 225, 162, x + 340, 162), fill="#CBD2D9", width=3)
        x += 390
    d.text((110, 228), "Request Category", font=f14, fill="#52606D")
    rounded(d, (110, 262, 515, 322), 10, "white", "#CBD2D9", 2)
    d.text((132, 278), "Other", font=f18, fill="#1F2933")
    d.text((560, 228), "Summary Title", font=f14, fill="#52606D")
    rounded(d, (560, 262, 1288, 322), 10, "white", "#CBD2D9", 2)
    d.text((582, 278), "Bug: alerts do not open", font=f18, fill="#1F2933")
    d.text((110, 352), "Detailed Description", font=f14, fill="#52606D")
    rounded(d, (110, 385, 1288, 445), 10, "white", "#CBD2D9", 2)
    d.text((132, 400), "What happened, what you expected, and a screenshot", font=f18, fill="#52606D")
    im.save(path)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, val in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LINE, size=6):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), str(size))
        tag.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_font(run, size=9.3, bold=False, color=INK, name="Aptos"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_text(doc, text, size=9.3, bold=False, color=INK, before=0, after=4, line=1.05, style=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    r = p.add_run(text)
    set_font(r, size, bold, color)
    return p


def add_bullet(doc, text, size=9.0, after=1.5):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.first_line_indent = Inches(-0.14)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.0
    set_font(p.add_run(text), size)
    return p


def heading(doc, text, level=1, before=5, after=3):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_font(r, 13.2 if level == 1 else 10.4, True, "000000")
    return p


def add_step(doc, number, title, body):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2.5)
    p.paragraph_format.line_spacing = 1.02
    r = p.add_run(f"{number}  {title}  ")
    set_font(r, 9.3, True, GREEN)
    set_font(p.add_run(body), 9.3, False, INK)


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(3)
    set_font(p.add_run(text), 7.8, False, MID)


def add_event_table(doc):
    rows = [
        ("You submit a maintenance request", "SiraNaBa records it, checks its urgency, and shows it in Maintenance."),
        ("Staff changes the ticket status", "A Maintenance alert appears so you can follow the repair."),
        ("A bill or payment changes", "A Payments alert reports the due date, receipt, or failed charge."),
        ("Management posts a facility notice", "A Community alert appears in the Notification Center."),
        ("You select an alert action", "The related ticket, bill, or notice opens and the alert is marked read."),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(2.55)
    table.columns[1].width = Inches(4.55)
    table.alignment = 1
    set_table_borders(table)
    hdr = table.rows[0].cells
    for i, label in enumerate(("Event", "Action you see")):
        set_cell_shading(hdr[i], GREEN)
        set_cell_margins(hdr[i], 90, 110, 90, 110)
        hdr[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = hdr[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(label), 8.8, True, WHITE)
    set_repeat_table_header(table.rows[0])
    for idx, (event, action) in enumerate(rows):
        cells = table.add_row().cells
        if idx % 2:
            for c in cells:
                set_cell_shading(c, "F4F8F5")
        for c in cells:
            set_cell_margins(c, 70, 100, 70, 100)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for c, value in zip(cells, (event, action)):
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            set_font(p.add_run(value), 8.2, False, INK)
    return table


def add_faq_table(doc):
    rows = [
        ("Why did my event not trigger an alert?", "Refresh the page, check Alert Settings, and confirm the correct category is on. New data may take about 15 seconds to appear. High-priority facility alerts cannot be muted."),
        ("Why can I not sign in?", "Check the email address and password, then try again. If the message says the details are incorrect, ask the administrator to confirm your account."),
        ("Why are there no alerts?", "Choose All Alerts, clear the search box, and select Retry if the page says notifications could not load."),
        ("Why did I still receive an alert after muting?", "SiraNaBa always sends high-priority facility alerts for safety and service disruptions."),
        ("Why is my ticket severity still loading?", "The system is assessing the request. Wait briefly, then refresh Maintenance. Report it if the status does not update."),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(2.38)
    table.columns[1].width = Inches(4.72)
    set_table_borders(table)
    for i, label in enumerate(("Question", "What to do")):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, "E5E7EB")
        set_cell_margins(cell, 80, 100, 80, 100)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(label), 8.6, True, "000000")
    set_repeat_table_header(table.rows[0])
    for idx, (q, a) in enumerate(rows):
        cells = table.add_row().cells
        if idx % 2:
            for c in cells:
                set_cell_shading(c, "FAFAFA")
        for c in cells:
            set_cell_margins(c, 62, 90, 62, 90)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for c, value in zip(cells, (q, a)):
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 0.98
            set_font(p.add_run(value), 7.9, q == value, INK)
    return table


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("SiraNaBa User Manual   |   ")
    set_font(run, 8, False, MID)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)


def build():
    settings_img = ASSET_DIR / "alert_settings_wireframe.png"
    bug_img = ASSET_DIR / "bug_report_wireframe.png"
    create_settings_wireframe(settings_img)
    create_bug_wireframe(bug_img)

    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.27)
    sec.page_height = Inches(11.69)
    sec.top_margin = Inches(0.48)
    sec.bottom_margin = Inches(0.45)
    sec.left_margin = Inches(0.58)
    sec.right_margin = Inches(0.58)
    sec.header_distance = Inches(0.2)
    sec.footer_distance = Inches(0.25)
    add_page_number(sec.footer.paragraphs[0])

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.3)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(3)

    for sname in ("Title", "Heading 1", "Heading 2"):
        st = styles[sname]
        st.font.name = "Aptos Display" if sname == "Title" else "Aptos"
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.font.bold = True

    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_after = Pt(1)
    set_font(title.add_run("SiraNaBa Notification User Manual"), 22, True, "000000", "Aptos Display")
    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(5)
    set_font(sub.add_run("Small team guide for tenants and staff"), 10.5, True, GREEN)
    add_text(
        doc,
        "Use this guide to sign in, choose the SiraNaBa alerts you want, quiet non-essential messages, understand what happens after an event, and report a problem. It is written for a small team of 5 to 20 users.",
        9.3, False, INK, after=4, line=1.03,
    )

    heading(doc, "Getting Started", 1, 3, 2)
    add_step(doc, "1", "Open SiraNaBa", "Use the web address provided by your administrator. No installation is needed on a shared computer; on a phone, open the same address in your browser.")
    add_step(doc, "2", "Sign in", "Enter your registered email address and password, then select Sign In. If you are using a shared device, sign out when finished.")
    add_step(doc, "3", "Open Notifications", "Select Notifications from the left menu. On a small screen, open the menu first, then select Notifications.")

    heading(doc, "Subscribe to the Alerts You Need", 1, 5, 2)
    add_step(doc, "1", "Open Alert Settings", "In the Notification Center, select the green Alert Settings button.")
    add_step(doc, "2", "Choose delivery methods", "Turn on Email, SMS, or Push Notifications. Green means on; gray means off.")
    add_step(doc, "3", "Choose event categories", "Turn on Payments for due dates and receipts, Maintenance for ticket updates and scheduled work, and Community for facility notices.")
    add_step(doc, "4", "Save", "Select Save Preferences. Keep at least one delivery method on if you want alerts outside the portal.")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(str(settings_img), width=Inches(6.85))
    add_caption(doc, "Wireframe based on the SiraNaBa Alert Settings screen")

    heading(doc, "How Event and Action Works", 1, 3, 2)
    add_text(doc, "An event is something that happens in SiraNaBa. The action is what the system does next for you.", 8.8, after=3)
    add_event_table(doc)

    doc.add_page_break()

    heading(doc, "Mute Notifications", 1, 0, 2)
    add_step(doc, "1", "Open Alert Settings", "Go to Notifications and select Alert Settings.")
    add_step(doc, "2", "Mute one type", "Turn off a category such as Community. Payments and Maintenance can stay on.")
    add_step(doc, "3", "Mute one delivery method", "Turn off Email, SMS, or Push while leaving the other methods on.")
    add_step(doc, "4", "Mute most alerts", "Turn off all categories and delivery methods, then select Save Preferences.")
    add_text(doc, "Important: high-priority facility alerts are always sent for safety and major service interruptions.", 8.8, True, GREEN, before=1, after=4)

    heading(doc, "Report a Bug", 1, 4, 2)
    add_text(doc, "A bug is when the system behaves differently from what you expected, such as a button not opening or an alert not appearing.", 8.8, after=3)
    add_step(doc, "1", "Open Submit Request", "Select Maintenance, then Submit Maintenance Request.")
    add_step(doc, "2", "Choose Other", "Use a title beginning with Bug, for example Bug: alerts do not open.")
    add_step(doc, "3", "Describe the problem", "Include what you selected, what happened, what you expected, the date and time, and the device or browser you used.")
    add_step(doc, "4", "Add evidence and submit", "Attach a screenshot if possible. Enter your location as your unit or work area, review the details, and select Submit Request.")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(str(bug_img), width=Inches(6.85))
    add_caption(doc, "Wireframe showing the recommended bug report details")

    heading(doc, "Quick Troubleshooting", 1, 3, 2)
    add_faq_table(doc)

    heading(doc, "Good Small Team Practice", 1, 4, 2)
    add_text(doc, "Use one shared rule: keep Maintenance alerts on for everyone responsible for repairs, keep Payments alerts on for the person handling accounts, and keep Community alerts on for team members who manage facility notices. Mark alerts as read after acting on them so the team can quickly see what still needs attention.", 8.6, after=0, line=1.0)

    core = doc.core_properties
    core.title = "SiraNaBa Notification User Manual"
    core.subject = "End-user documentation for notification settings and bug reporting"
    core.author = "SiraNaBa Project"
    core.keywords = "SiraNaBa, user manual, notifications, small team"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
